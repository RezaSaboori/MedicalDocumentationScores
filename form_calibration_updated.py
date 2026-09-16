from __future__ import annotations

from bisect import bisect_left, bisect_right
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

INPUT_FILE = Path(
    "MedicalDocsStatus_2026-09-16_14_16_08.xlsx"
)

OUTPUT_FILE = Path(
    "MedicalDocsStatus_2026-09-16_14_16_08_calibrated.xlsx"
)

SHEET_NAME = "Medical Docs Status"

# Reliability constant:
#
# lambda = N / (N + K)
#
# Larger K -> more conservative calibration.
K = 100

# Minimum reliability required when a raw class 3 or 4 form
# is promoted to calibrated class 5 by relative population evidence.
#
# With K = 100:
#
#     reliability = N / (N + 100)
#
# reliability >= 0.50 corresponds to N >= 100.
#
# IMPORTANT:
# Raw class 5 does NOT require this minimum reliability to remain
# calibrated class 5 when its calibrated score itself remains > 95.
MIN_RELATIVE_CLASS5_RELIABILITY = 0.50


# ============================================================
# COLUMN NAMES
# ============================================================

FORM_COLUMN = "پرونده الکترونیک"
RAW_SCORE_COLUMN = "امتیاز خام"

CALIBRATED_SCORE_COLUMN = "امتیاز کالیبره"
RAW_CLASS_COLUMN = "کلاس امتیاز خام"
CALIBRATED_CLASS_COLUMN = "کلاس امتیاز کالیبره"
COMBINED_STATUS_COLUMN = "وضعیت ترکیبی"


# ============================================================
# CLASS LABELS
# ============================================================

CLASS_FA = {
    0: "خالی",
    1: "کمتر از حداقل انتظار",
    2: "حداقل قابل قبول",
    3: "سطح قابل قبول",
    4: "خوب",
    5: "فراتر از انتظار",
}


# ============================================================
# HELPERS
# ============================================================

def to_float(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

    try:
        return float(value)

    except (TypeError, ValueError):
        return None


def round_2(value: float) -> float:
    """
    Round using conventional ROUND_HALF_UP rather than
    Python's banker's rounding.
    """

    return float(
        Decimal(str(value)).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        )
    )


def score_to_units(score: float) -> int:
    """
    Convert score into integer hundredths.

    Example:
        57.32 -> 5732

    This avoids floating-point equality problems when detecting
    tied raw scores.
    """

    return int(
        (
            Decimal(str(score))
            * Decimal("100")
        ).quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )


def normalize_raw_class(
    raw_score,
    existing_class,
):
    """
    Existing raw class is authoritative because class 5 may
    depend on information from the original scoring engine
    that is not reconstructable from raw_score alone.

    If the class is missing, reconstruct only classes 0..4.
    """

    try:
        if existing_class is not None:
            value = int(
                float(existing_class)
            )

            if 0 <= value <= 5:
                return value

    except (TypeError, ValueError):
        pass

    # No raw score means truly empty.
    if raw_score is None:
        return 0

    if raw_score <= 25.00:
        return 1

    if raw_score <= 50.00:
        return 2

    if raw_score <= 75.00:
        return 3

    # We deliberately do NOT infer raw class 5 here because
    # class 5 may require trophy/capping information.
    return 4


# ============================================================
# MID-RANK PERCENTILE
# ============================================================

def midrank_percentile(
    x_units: int,
    sorted_reference_units: list[int],
) -> float:
    """
    P = 100 * [N(X < x) + 0.5*N(X = x)] / N
    """

    n = len(
        sorted_reference_units
    )

    if n == 0:
        return 0.0

    lower = bisect_left(
        sorted_reference_units,
        x_units,
    )

    upper = bisect_right(
        sorted_reference_units,
        x_units,
    )

    equal = upper - lower

    percentile = (
        100.0
        * (
            lower
            + 0.5 * equal
        )
        / n
    )

    return percentile


# ============================================================
# CALIBRATED CLASS
# ============================================================

def calibration_class(
    calibration_score: float,
    raw_class: int,
    percentile: float,
    reliability: float,
) -> int:
    """
    Final calibrated classification.

    Class 5 has two valid pathways:

    Path A — absolute excellence:
        calibration_score > 95
        raw_class == 5

        No minimum reference size/reliability is required because
        the form was already exceptional under the absolute raw
        scoring system.

    Path B — relative promotion:
        calibration_score > 95
        raw_class >= 3
        percentile >= 95
        reliability >= MIN_RELATIVE_CLASS5_RELIABILITY

        This allows a raw class 3 or 4 form to be promoted only
        when the form-specific population evidence is sufficiently
        strong.

    Raw classes 1 and 2 can never become calibrated class 5 because
    of the absolute-quality ceilings applied before this function.
    """

    if raw_class == 0:
        return 0

    if calibration_score <= 25.00:
        return 1

    if calibration_score <= 50.00:
        return 2

    if calibration_score <= 75.00:
        return 3

    if calibration_score <= 95.00:
        return 4

    # --------------------------------------------------------
    # Path A: absolute excellence
    # --------------------------------------------------------

    if raw_class == 5:
        return 5

    # --------------------------------------------------------
    # Path B: statistically reliable relative promotion
    # --------------------------------------------------------

    if (
        raw_class >= 3
        and percentile >= 95.00
        and reliability >= MIN_RELATIVE_CLASS5_RELIABILITY
    ):
        return 5

    return 4


# ============================================================
# MAIN
# ============================================================

workbook = load_workbook(
    INPUT_FILE
)

worksheet = workbook[
    SHEET_NAME
]


# ------------------------------------------------------------
# Detect columns from headers
# ------------------------------------------------------------

header_to_column = {}

for cell in worksheet[1]:
    if cell.value is not None:
        header_to_column[
            str(cell.value).strip()
        ] = cell.column


required_columns = [
    FORM_COLUMN,
    RAW_SCORE_COLUMN,
    CALIBRATED_SCORE_COLUMN,
    RAW_CLASS_COLUMN,
    CALIBRATED_CLASS_COLUMN,
    COMBINED_STATUS_COLUMN,
]

missing_columns = [
    column
    for column in required_columns
    if column not in header_to_column
]

if missing_columns:
    raise ValueError(
        "Missing required columns: "
        + ", ".join(
            missing_columns
        )
    )


form_col = header_to_column[
    FORM_COLUMN
]

raw_score_col = header_to_column[
    RAW_SCORE_COLUMN
]

calibrated_score_col = header_to_column[
    CALIBRATED_SCORE_COLUMN
]

raw_class_col = header_to_column[
    RAW_CLASS_COLUMN
]

calibrated_class_col = header_to_column[
    CALIBRATED_CLASS_COLUMN
]

combined_status_col = header_to_column[
    COMBINED_STATUS_COLUMN
]


# ============================================================
# PASS 1
#
# Build ONE-MONTH reference population separately for each form.
#
# Truly empty forms (raw class 0) are NOT included.
# ============================================================

reference_scores = defaultdict(
    list
)

row_information = {}


for row in range(
    2,
    worksheet.max_row + 1,
):

    form_type = worksheet.cell(
        row=row,
        column=form_col,
    ).value

    raw_score = to_float(
        worksheet.cell(
            row=row,
            column=raw_score_col,
        ).value
    )

    existing_raw_class = (
        worksheet.cell(
            row=row,
            column=raw_class_col,
        ).value
    )

    raw_class = normalize_raw_class(
        raw_score=raw_score,
        existing_class=existing_raw_class,
    )

    # Normalize raw-class column.
    worksheet.cell(
        row=row,
        column=raw_class_col,
    ).value = raw_class

    row_information[row] = {
        "form_type": form_type,
        "raw_score": raw_score,
        "raw_class": raw_class,
    }

    # Empty forms are excluded from calibration reference.
    if (
        form_type is None
        or raw_score is None
        or raw_class == 0
    ):
        continue

    form_key = str(
        form_type
    ).strip()

    reference_scores[
        form_key
    ].append(
        score_to_units(
            raw_score
        )
    )


# Sort once so percentile lookup becomes O(log N).
for form_key in reference_scores:
    reference_scores[
        form_key
    ].sort()


# ============================================================
# PASS 2
#
# Calculate calibrated score for every row.
# ============================================================

for row in range(
    2,
    worksheet.max_row + 1,
):

    info = row_information[
        row
    ]

    form_type = info[
        "form_type"
    ]

    raw_score = info[
        "raw_score"
    ]

    raw_class = info[
        "raw_class"
    ]

    calibrated_score_cell = (
        worksheet.cell(
            row=row,
            column=calibrated_score_col,
        )
    )

    calibrated_class_cell = (
        worksheet.cell(
            row=row,
            column=calibrated_class_col,
        )
    )

    combined_status_cell = (
        worksheet.cell(
            row=row,
            column=combined_status_col,
        )
    )

    # --------------------------------------------------------
    # Empty form
    # --------------------------------------------------------

    if (
        raw_class == 0
        or raw_score is None
        or form_type is None
    ):

        calibrated_score_cell.value = 0.0
        calibrated_class_cell.value = 0
        combined_status_cell.value = (
            CLASS_FA[0]
        )

        continue


    form_key = str(
        form_type
    ).strip()

    references = reference_scores.get(
        form_key,
        [],
    )

    reference_n = len(
        references
    )


    # --------------------------------------------------------
    # Relative percentile P
    # --------------------------------------------------------

    x_units = score_to_units(
        raw_score
    )

    percentile = midrank_percentile(
        x_units=x_units,
        sorted_reference_units=references,
    )


    # --------------------------------------------------------
    # Reliability
    #
    # lambda = N / (N + K)
    # --------------------------------------------------------

    reliability = (
        reference_n
        / (
            reference_n
            + K
        )
        if reference_n > 0
        else 0.0
    )


    # --------------------------------------------------------
    # Reliability-adjusted score
    #
    # S = (1-lambda)*R + lambda*P
    # --------------------------------------------------------

    adjusted_score = (
        (
            1.0
            - reliability
        )
        * raw_score
        +
        reliability
        * percentile
    )


    # --------------------------------------------------------
    # Absolute quality guardrail
    # --------------------------------------------------------

    class_ceiling = {
        0: 0.00,
        1: 75.00,
        2: 95.00,
        3: 100.00,
        4: 100.00,
        5: 100.00,
    }[
        raw_class
    ]


    calibrated_score = min(
        adjusted_score,
        class_ceiling,
    )


    # --------------------------------------------------------
    # Above-expectation guardrail
    #
    # Class 5 has TWO valid pathways:
    #
    # Path A — absolute excellence:
    #   raw_class == 5
    #   calibrated_score > 95
    #
    # Path B — relative promotion:
    #   raw_class >= 3
    #   calibrated_score > 95
    #   percentile >= 95
    #   reliability >= 0.50
    #
    # Therefore rare forms can preserve genuine raw class-5
    # performance, while class 3/4 promotion still requires
    # sufficiently strong population evidence.
    # --------------------------------------------------------

    absolute_excellence = (
        raw_class == 5
    )

    relative_excellence = (
        raw_class >= 3
        and percentile >= 95.00
        and reliability
        >= MIN_RELATIVE_CLASS5_RELIABILITY
    )

    class_5_eligible = (
        calibrated_score > 95.00
        and (
            absolute_excellence
            or relative_excellence
        )
    )

    if (
        calibrated_score > 95.00
        and not class_5_eligible
    ):
        calibrated_score = 95.00


    # --------------------------------------------------------
    # Round final score
    # --------------------------------------------------------

    calibrated_score = round_2(
        calibrated_score
    )


    # --------------------------------------------------------
    # Final calibrated class
    # --------------------------------------------------------

    calibrated_class = (
        calibration_class(
            calibration_score=(
                calibrated_score
            ),
            raw_class=raw_class,
            percentile=percentile,
            reliability=reliability,
        )
    )


    # --------------------------------------------------------
    # Write requested columns
    # --------------------------------------------------------

    calibrated_score_cell.value = (
        calibrated_score
    )

    calibrated_class_cell.value = (
        calibrated_class
    )

    combined_status_cell.value = (
        CLASS_FA[
            calibrated_class
        ]
    )


# ============================================================
# FORMATTING
# ============================================================

for row in range(
    2,
    worksheet.max_row + 1,
):
    worksheet.cell(
        row=row,
        column=calibrated_score_col,
    ).number_format = "0.00"

    worksheet.cell(
        row=row,
        column=raw_class_col,
    ).number_format = "0"

    worksheet.cell(
        row=row,
        column=calibrated_class_col,
    ).number_format = "0"


# ============================================================
# SAVE
# ============================================================

workbook.save(
    OUTPUT_FILE
)

print(
    f"Saved calibrated workbook: {OUTPUT_FILE}"
)